#!/usr/bin/env python3
"""layout_ocr_to_docx: rebuild an image-based PDF (or image) as an editable DOCX
that keeps the page layout. Runs 100% offline.

For every page it creates, at the original position and size:
  * text boxes with real, copyable text (font size and colour estimated,
    bold detected, font name chosen by you with --font)
  * pictures for the non-text regions (photos, charts, logos)
  * filled rectangles (coloured bars/boxes) and thin line shapes
    (box borders, table grids, underlines)

Examples:
  python layout_ocr_to_docx.py scan.pdf
  python layout_ocr_to_docx.py scan.pdf -o result.docx --font "Times New Roman"
  python layout_ocr_to_docx.py scan.pdf --pages 1-3 --debug
"""
import argparse
import io
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

import cv2
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Emu, Pt
from lxml import etree
from PIL import Image, ImageOps

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
EMU_PER_INCH = 914400
NSDECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)


# ---------------------------------------------------------------- helpers
def parse_pages(spec, total):
    if not spec:
        return list(range(1, total + 1))
    pages = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            pages.extend(range(int(a) if a.strip() else 1,
                               (int(b) if b.strip() else total) + 1))
        else:
            pages.append(int(part))
    out = []
    for p in pages:
        if 1 <= p <= total and p not in out:
            out.append(p)
    if not out:
        sys.exit(f"No valid pages in '{spec}' (document has {total} pages).")
    return out


def load_pages(path, a):
    """Yield (PIL image, dpi) for each page, batching Poppler renders to bound memory and maximize speed."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pdf2image import convert_from_path, pdfinfo_from_path
        total = pdfinfo_from_path(str(path), poppler_path=a.poppler_path)["Pages"]
        pages = parse_pages(a.pages, total)
        batch_size = max(1, getattr(a, "page_batch", 5))
        for i in range(0, len(pages), batch_size):
            chunk = pages[i:i + batch_size]
            if chunk == list(range(chunk[0], chunk[-1] + 1)):
                imgs = convert_from_path(str(path), dpi=a.dpi, first_page=chunk[0],
                                         last_page=chunk[-1], poppler_path=a.poppler_path)
                for img in imgs:
                    yield img.convert("RGB"), a.dpi
            else:
                for n in chunk:
                    img = convert_from_path(str(path), dpi=a.dpi, first_page=n, last_page=n,
                                            poppler_path=a.poppler_path)[0]
                    yield img.convert("RGB"), a.dpi
    elif suffix in IMAGE_EXTS:
        img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
        yield img, img.width / a.image_page_width
    else:
        sys.exit(f"Unsupported file type: {path}")


def snap_color(c):
    c = tuple(int(v) for v in c)
    spread = max(c) - min(c)
    # Only snap true dark neutrals to pure black, and true light neutrals to pure white.
    # Preserve saturated or tinted colors (e.g. dark blue, green, cream, pastel).
    if max(c) < 36 and spread < 16:
        c = (0, 0, 0)
    elif min(c) > 244 and spread < 12:
        c = (255, 255, 255)
    return "%02X%02X%02X" % c


def detect_word_italic(bin_img):
    """Detect whether a binary ink crop (1=ink, 0=bg) is italic using shear projection variance."""
    h, w = bin_img.shape
    if w < 10 or h < 8:
        return False
    v0 = float(bin_img.sum(axis=0).var())
    if v0 <= 0:
        return False
    best_v, best_ang = v0, 0
    for ang in (6, 11, 16):
        tan_a = np.tan(np.radians(ang))
        M = np.float32([[1, tan_a, 0], [0, 1, 0]])
        sheared = cv2.warpAffine(bin_img, M, (w + int(h * 0.35), h))
        v = float(sheared.sum(axis=0).var())
        if v > best_v:
            best_v, best_ang = v, ang
    return (best_ang >= 6) and (best_v > 1.05 * v0)


def clean_text(t):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", t)


# ------------------------------------------------------------------- OCR
def _ocr_pass(tess, img, a, psm):
    cfg = f"--psm {psm}" if psm else ""
    d = tess.image_to_data(img, lang=a.lang, config=cfg, output_type=tess.Output.DICT)
    words = []
    for i, t in enumerate(d["text"]):
        t = clean_text(t or "").strip()
        if not t:
            continue
        try:
            conf = float(d["conf"][i])
        except (TypeError, ValueError):
            conf = -1
        if conf < a.min_conf or d["width"][i] < 2 or d["height"][i] < 4:
            continue
        words.append(dict(text=t, x=int(d["left"][i]), y=int(d["top"][i]),
                          w=int(d["width"][i]), h=int(d["height"][i]),
                          key=(d["block_num"][i], d["par_num"][i], d["line_num"][i])))
    return words


def _overlaps(e, w):
    ix = max(0, min(e["x"] + e["w"], w["x"] + w["w"]) - max(e["x"], w["x"]))
    iy = max(0, min(e["y"] + e["h"], w["y"] + w["h"]) - max(e["y"], w["y"]))
    return ix * iy > 0.25 * min(e["w"] * e["h"], w["w"] * w["h"])


def ocr_words(tess, img, a):
    words = _ocr_pass(tess, img, a, a.psm)
    if not a.no_sparse_pass:
        # second pass finds isolated numbers/labels the normal pass skips
        for e in _ocr_pass(tess, img, a, 11):
            if not any(_overlaps(e, w) for w in words):
                e["key"] = (10000 + e["key"][0], e["key"][1], e["key"][2])
                words.append(e)
    return words


def page_background(arr):
    small = arr[::8, ::8].reshape(-1, 3)
    q = small // 16
    keys = q[:, 0].astype(np.int32) * 256 + q[:, 1] * 16 + q[:, 2]
    vals, counts = np.unique(keys, return_counts=True)
    sel = keys == vals[counts.argmax()]
    return tuple(int(v) for v in np.median(small[sel], axis=0))


def word_style(arr, wd, pad=3, check_italic=True):
    """Return (text colour, stroke width in px, local background colour, is_italic)."""
    H, W = arr.shape[:2]
    x0, y0 = max(0, wd["x"] - pad), max(0, wd["y"] - pad)
    x1, y1 = min(W, wd["x"] + wd["w"] + pad), min(H, wd["y"] + wd["h"] + pad)
    crop = arr[y0:y1, x0:x1].astype(np.int16)
    if crop.shape[0] < 3 or crop.shape[1] < 3:
        return (0, 0, 0), 0.0, (255, 255, 255), False
    ring = np.concatenate([crop[0], crop[-1], crop[:, 0], crop[:, -1]])
    bg = np.median(ring, axis=0)
    dist = np.abs(crop - bg).sum(axis=2)
    d_max = dist.max()
    ink_thresh = min(100, max(30, int(0.35 * d_max))) if d_max > 0 else 100
    ink = dist >= ink_thresh
    if ink.sum() < 4:
        col = (255, 255, 255) if bg.mean() < 128 else (0, 0, 0)
        return col, 0.0, tuple(int(v) for v in bg), False
    sel = ink & (dist >= np.percentile(dist[ink], 60))
    col = np.median(crop[sel], axis=0)
    strict = dist > 0.5 * d_max
    if strict.sum() > 0:
        dt = cv2.distanceTransform(strict.astype(np.uint8), cv2.DIST_L2, 3)
        stroke = 4.0 * float(dt[strict].mean())
    else:
        stroke = 0.0
    is_ita = False
    if check_italic and wd["w"] >= 10 and wd["h"] >= 8 and strict.sum() >= 15:
        is_ita = detect_word_italic(strict.astype(np.uint8))
    return tuple(int(v) for v in col), stroke, tuple(int(v) for v in bg), is_ita


# ------------------------------------------------- graphics (pictures etc)
def snap_table_grid(lines, tolerance=3):
    """Snap nearly-touching line endpoints to form continuous table grids."""
    h_lines = [l for l in lines if l["w"] > l["h"]]
    v_lines = [l for l in lines if l["h"] >= l["w"]]
    for hl in h_lines:
        hx0, hx1, hy0, hy1 = hl["x"], hl["x"] + hl["w"], hl["y"], hl["y"] + hl["h"]
        for vl in v_lines:
            vx0, vx1, vy0, vy1 = vl["x"], vl["x"] + vl["w"], vl["y"], vl["y"] + vl["h"]
            if abs(hx0 - vx0) <= tolerance and (vy0 - tolerance <= hy1 and vy1 + tolerance >= hy0):
                hl["w"] += hl["x"] - vx0
                hl["x"] = vx0
            if abs(hx1 - vx1) <= tolerance and (vy0 - tolerance <= hy1 and vy1 + tolerance >= hy0):
                hl["w"] = vx1 - hl["x"]
            if abs(vy0 - hy0) <= tolerance and (hx0 - tolerance <= vx1 and hx1 + tolerance >= vx0):
                vl["h"] += vl["y"] - hy0
                vl["y"] = hy0
            if abs(vy1 - hy1) <= tolerance and (hx0 - tolerance <= vx1 and hx1 + tolerance >= vx0):
                vl["h"] = hy1 - vl["y"]
    return lines


def match_underlines(lines, words, a):
    """Detect horizontal lines that serve as underlines for text runs."""
    if getattr(a, "no_underline", False):
        return lines
    remaining_lines = []
    for ln in lines:
        if ln["w"] <= 2.5 * ln["h"] or ln["h"] > 16:
            remaining_lines.append(ln)
            continue
        matched = []
        for wd in words:
            y_diff = ln["y"] - (wd["y"] + wd["h"])
            if -10 <= y_diff <= max(25, int(0.65 * wd["h"])):
                if (wd["x"] + wd["w"] >= ln["x"] - 8) and (wd["x"] <= ln["x"] + ln["w"] + 8):
                    matched.append(wd)
        if matched:
            span_x0 = min(w["x"] for w in matched)
            span_x1 = max(w["x"] + w["w"] for w in matched)
            span_w = span_x1 - span_x0
            if span_w >= 0.65 * ln["w"] and ln["w"] <= 1.45 * span_w:
                for wd in matched:
                    wd["underline"] = True
                continue
        remaining_lines.append(ln)
    return remaining_lines


def detect_graphics(arr, words, bg, a):
    H, W = arr.shape[:2]
    work = arr.copy()
    for wd in words:  # paint over text with proportional margin so edge residues don't linger
        pad_x = max(4, int(0.04 * wd["w"]))
        pad_y = max(5, int(0.12 * wd["h"]))
        work[max(0, wd["y"] - pad_y):min(H, wd["y"] + wd["h"] + pad_y),
             max(0, wd["x"] - pad_x):min(W, wd["x"] + wd["w"] + pad_x)] = wd["bg"]
    dist = np.abs(work.astype(np.int16) - np.array(bg, dtype=np.int16)).sum(axis=2)
    mask = ((dist > a.bg_tol) * 255).astype(np.uint8)

    fills, images = [], []
    page_area = W * H
    min_img = a.min_image_pct / 100.0 * page_area
    min_fill = 0.0002 * page_area

    # 1) thick solid blobs: flat colour -> filled shape, busy -> picture
    if not a.no_images:
        solid = cv2.morphologyEx(mask, cv2.MORPH_OPEN, np.ones((9, 9), np.uint8))
        kc = max(9, int(0.01 * W))
        solid = cv2.morphologyEx(solid, cv2.MORPH_CLOSE, np.ones((kc, kc), np.uint8))
        for c in cv2.findContours(solid, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0]:
            x, y, w, h = cv2.boundingRect(c)
            m = mask[y:y + h, x:x + w] > 0
            if m.sum() == 0:
                continue
            reg = work[y:y + h, x:x + w]
            dens = float(m.mean())
            # Erode mask to evaluate interior fill without border interference
            m_uint8 = m.astype(np.uint8)
            m_inner = cv2.erode(m_uint8, np.ones((5, 5), np.uint8)) > 0
            eval_m = m_inner if m_inner.sum() >= 0.25 * m.sum() else m
            std = float(reg[eval_m].std(axis=0).mean())
            if std < 18 and dens > 0.75 and w * h >= min_fill:
                fills.append(dict(x=x, y=y, w=w, h=h,
                                  color=snap_color(np.median(reg[eval_m], axis=0))))
            elif w * h >= min_img and dens > 0.25:
                images.append(dict(x=x, y=y, w=w, h=h))

    resid = mask.copy()
    for r in fills + images:
        resid[r["y"]:r["y"] + r["h"], r["x"]:r["x"] + r["w"]] = 0

    # 2) long straight lines (box borders, table grids, underlines)
    L = max(25, int(0.02 * W))
    hl = cv2.morphologyEx(resid, cv2.MORPH_OPEN, np.ones((1, L), np.uint8))
    vl = cv2.morphologyEx(resid, cv2.MORPH_OPEN, np.ones((L, 1), np.uint8))

    # 3) thin-stroke graphics (charts, logos, line art) -> pictures
    if not a.no_images:
        rest = cv2.bitwise_and(resid, cv2.bitwise_not(
            cv2.dilate(cv2.bitwise_or(hl, vl), np.ones((3, 3), np.uint8))))
        k2 = max(15, int(0.02 * W))
        rest = cv2.morphologyEx(rest, cv2.MORPH_CLOSE, np.ones((k2, k2), np.uint8))
        for c in cv2.findContours(rest, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0]:
            x, y, w, h = cv2.boundingRect(c)
            if w * h < min_img:
                continue
            if (resid[y:y + h, x:x + w] > 0).mean() < 0.02:
                continue
            overlaps = False
            for r in images:
                ix = max(0, min(x + w, r["x"] + r["w"]) - max(x, r["x"]))
                iy = max(0, min(y + h, r["y"] + r["h"]) - max(y, r["y"]))
                if ix * iy > 0.3 * min(w * h, r["w"] * r["h"]):
                    overlaps = True
                    break
            if not overlaps:
                images.append(dict(x=x, y=y, w=w, h=h))

    def inside_image(cx, cy):
        return any(r["x"] <= cx <= r["x"] + r["w"] and r["y"] <= cy <= r["y"] + r["h"]
                   for r in images)

    lines = []
    if not a.no_lines:
        thin = max(14, int(0.012 * W))
        for kind, lm in (("h", hl), ("v", vl)):
            n, lab, stats, _ = cv2.connectedComponentsWithStats(lm, connectivity=8)
            for i in range(1, n):
                x, y, w, h = (int(v) for v in stats[i][:4])
                long_, short = (w, h) if kind == "h" else (h, w)
                if long_ < L or short > thin or inside_image(x + w // 2, y + h // 2):
                    continue
                sel = lab[y:y + h, x:x + w] == i
                color = snap_color(np.median(work[y:y + h, x:x + w][sel], axis=0))
                lines.append(dict(x=x, y=y, w=w, h=h, color=color))
        lines = snap_table_grid(lines)
    left = mask.copy()
    for r in fills + images:
        left[r["y"]:r["y"] + r["h"], r["x"]:r["x"] + r["w"]] = 0
    left[cv2.dilate(cv2.bitwise_or(hl, vl), np.ones((5, 5), np.uint8)) > 0] = 0
    return fills, images, lines, left


def orphan_words(tess, img, arr, left, images, a):
    """Read leftover ink blobs (e.g. lone digits) one by one with Tesseract."""
    H, W = arr.shape[:2]
    kx = max(3, int(0.006 * W))
    grp = cv2.dilate(left, np.ones((5, kx), np.uint8))
    cands = []
    for c in cv2.findContours(grp, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0]:
        x, y, w, h = cv2.boundingRect(c)
        ys, xs = np.nonzero(left[y:y + h, x:x + w])
        if len(xs) == 0:
            continue
        x, y = x + int(xs.min()), y + int(ys.min())
        w, h = int(xs.max() - xs.min() + 1), int(ys.max() - ys.min() + 1)
        if h < 14 or h > 160 or w < 5 or w > h * 14:
            continue
        cx, cy = x + w / 2, y + h / 2
        if any(r["x"] <= cx <= r["x"] + r["w"] and r["y"] <= cy <= r["y"] + r["h"] for r in images):
            continue
        cands.append((y, x, w, h))
    out = []
    for n, (y, x, w, h) in enumerate(sorted(cands)[:150]):
        pad = 12
        crop = Image.fromarray(arr[max(0, y - pad):y + h + pad, max(0, x - pad):x + w + pad])
        d = tess.image_to_data(crop, lang=a.lang, config="--psm 7", output_type=tess.Output.DICT)
        toks = [(clean_text(t).strip(), float(cf)) for t, cf in zip(d["text"], d["conf"])
                if clean_text(t or "").strip() and float(cf) >= 0]
        if not toks:
            continue
        text = " ".join(t for t, _ in toks)
        conf = sum(cf for _, cf in toks) / len(toks)
        if conf >= max(a.min_conf, 55.0) and len(text) <= 15:
            if len(text) == 1 and not (text.isalnum() and conf >= 75.0):
                continue
            if any(ch.isalnum() for ch in text):
                out.append(dict(text=text, x=x, y=y, w=w, h=h, key=(20000 + n, 1, 1), is_orphan=True))
    return out


# ----------------------------------------------------------- text boxes
def make_text_boxes(words, dpi, a, page_w=None):
    px2tw = 1440.0 / dpi
    y_off = a.y_offset * dpi / 72.0

    by_line = {}
    for wd in words:
        by_line.setdefault(wd["key"], []).append(wd)

    raw = {}
    rels = {}
    valid_rels = []
    for key, ws in by_line.items():
        has_ascenders = any(any(c.isupper() or c in "bdfhklt" for c in wd["text"]) for wd in ws)
        divisor = 0.76 if has_ascenders else 0.54
        size = float(np.percentile([d["h"] for d in ws], 80)) / divisor
        raw[key] = size
        strk = float(np.median([d["stroke"] for d in ws])) / max(size, 1.0)
        rels[key] = strk
        if not any(wd.get("is_orphan", False) for wd in ws):
            valid_rels.append(strk)
    page_med = float(np.median(valid_rels)) if valid_rels else (float(np.median(list(rels.values()))) if rels else 0.0)

    # snap similar sizes together so a paragraph keeps one font size
    snapped, cluster = {}, []

    def close_cluster():
        val = float(np.percentile([raw[k] for k in cluster], 75))
        snapped.update({k: val for k in cluster})

    for key in sorted(raw, key=raw.get):
        if cluster and raw[key] > raw[cluster[0]] * 1.12:
            close_cluster()
            cluster = []
        cluster.append(key)
    if cluster:
        close_cluster()

    line_objs = []
    for key, ws in by_line.items():
        ws.sort(key=lambda d: d["x"])
        size_px = snapped[key]
        baseline = float(np.percentile([d["y"] + d["h"] for d in ws], 30))
        bold = ((not a.no_bold) and len(valid_rels) > 1 and
                rels[key] > max(a.bold_factor * page_med, 0.105))
        med = np.median([d["color"] for d in ws], axis=0)
        for d in ws:
            if np.abs(np.array(d["color"]) - med).sum() < 90:
                d["color"] = tuple(int(v) for v in med)
        segs = [[ws[0]]]
        for prev, cur in zip(ws, ws[1:]):
            if cur["x"] - (prev["x"] + prev["w"]) > 1.5 * size_px:
                segs.append([cur])
            else:
                segs[-1].append(cur)
        for si, seg in enumerate(segs):
            line_objs.append(dict(par=key[:2], line=key[2], col=si, words=seg,
                                  size_px=size_px, baseline=baseline, bold=bold,
                                  x0=min(d["x"] for d in seg),
                                  x1=max(d["x"] + d["w"] for d in seg)))

    groups = {}
    for lo in line_objs:
        gk = (lo["par"], lo["line"], lo["col"]) if a.box_per_line else (lo["par"], lo["col"])
        groups.setdefault(gk, []).append(lo)

    boxes = []
    no_align = getattr(a, "no_align", False)
    for glines in groups.values():
        glines.sort(key=lambda l: l["baseline"])
        size_px = float(np.median([l["size_px"] for l in glines]))
        size_pt = max(4.0, round(size_px * 72.0 / dpi * a.font_scale * 2) / 2)
        if len(glines) > 1:
            diffs = np.diff([l["baseline"] for l in glines])
            pitch = float(np.median(diffs))
        else:
            pitch = size_px * 1.2
        pitch = max(pitch, size_px * 1.0)
        x_left = min(l["x0"] for l in glines)
        y_top = glines[0]["baseline"] - (pitch - 0.21 * size_px) + y_off

        # Alignment detection
        align = "left"
        if not no_align and page_w:
            if len(glines) == 1:
                cx = (glines[0]["x0"] + glines[0]["x1"]) / 2.0
                if abs(cx - page_w / 2.0) < 0.04 * page_w:
                    align = "center"
                elif glines[0]["x1"] > 0.84 * page_w and glines[0]["x0"] > 0.35 * page_w:
                    align = "right"
            else:
                cxs = [(l["x0"] + l["x1"]) / 2.0 for l in glines]
                lefts = [l["x0"] for l in glines]
                rights = [l["x1"] for l in glines]
                if float(np.std(cxs)) < 15 and float(np.std(lefts)) > 20:
                    align = "center"
                elif float(np.std(rights)) < 15 and float(np.std(lefts)) > 20:
                    align = "right"
                elif len(glines) >= 3 and float(np.std(lefts)) < 15 and float(np.std(rights[:-1])) < 20:
                    align = "both"

        paras, total_h = [], 0.0
        prev_base = None
        for lo in glines:
            before = 0.0
            if prev_base is not None:
                extra = (lo["baseline"] - prev_base) - pitch
                before = extra if extra > 0.2 * pitch else 0.0
            prev_base = lo["baseline"]
            total_h += before + pitch
            runs, cur = [], None
            for i, wd in enumerate(lo["words"]):
                col = snap_color(wd["color"])
                txt = wd["text"] + (" " if i < len(lo["words"]) - 1 else "")
                ita = wd.get("italic", False)
                u = wd.get("underline", False)
                if cur and cur[1] == col and cur[2] == ita and cur[3] == u:
                    cur[0] += txt
                else:
                    cur = [txt, col, ita, u]
                    runs.append(cur)
            half = int(round(size_pt * 2))
            font = escape(a.font, {'"': "&quot;"})
            b = "<w:b/>" if lo["bold"] else ""
            run_xml = "".join(
                f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" '
                f'w:cs="{font}" w:eastAsia="{font}"/>{b}'
                f'{"<w:i/>" if ita else ""}'
                f'{"<w:u w:val=\"single\"/>" if u else ""}'
                f'<w:color w:val="{c}"/>'
                f'<w:sz w:val="{half}"/><w:szCs w:val="{half}"/></w:rPr>'
                f'<w:t xml:space="preserve">{escape(t)}</w:t></w:r>' for t, c, ita, u in runs)
            align_xml = f'<w:jc w:val="{align}"/>' if align != "left" else ""
            ind_val = 0 if align in ("center", "right") else int((lo["x0"] - x_left) * px2tw)
            paras.append(
                f'<w:p><w:pPr>{align_xml}<w:spacing w:before="{int(before * px2tw)}" w:after="0" '
                f'w:line="{int(pitch * px2tw)}" w:lineRule="exact"/>'
                f'<w:ind w:left="{ind_val}"/></w:pPr>{run_xml}</w:p>')
        width = (max(l["x1"] for l in glines) - x_left) * 1.35 + size_px
        boxes.append(dict(x=x_left, y=y_top, w=width, h=total_h * 1.1 + size_px,
                          xml="".join(paras)))
    return boxes


# ------------------------------------------------------------- DOCX XML
class Emitter:
    def __init__(self, doc):
        self.doc = doc
        self.next_id = 1
        self.z = 1000

    def _anchor(self, name, x, y, w, h, graphic, dpi):
        e = lambda v: int(round(v * EMU_PER_INCH / dpi))
        self.next_id += 1
        self.z += 1
        return etree.fromstring(
            f'<w:r {NSDECL}><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" '
            f'simplePos="0" relativeHeight="{self.z}" behindDoc="0" locked="0" '
            f'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{e(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{e(y)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{max(e(w), 1)}" cy="{max(e(h), 1)}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{self.next_id}" name="{name} {self.next_id}"/>'
            f'<wp:cNvGraphicFramePr/>{graphic}</wp:anchor></w:drawing></w:r>')

    @staticmethod
    def _shape(cx, cy, fill, txbx=None):
        fill_xml = (f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
                    if fill else '<a:noFill/>')
        tx_attr = ' txBox="1"' if txbx else ""
        tx = f'<wps:txbx><w:txbxContent>{txbx}</w:txbxContent></wps:txbx>' if txbx else ""
        return (
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f'<wps:wsp><wps:cNvSpPr{tx_attr}/><wps:spPr><a:xfrm><a:off x="0" y="0"/>'
            f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'{fill_xml}<a:ln><a:noFill/></a:ln></wps:spPr>{tx}'
            '<wps:bodyPr rot="0" vert="horz" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" '
            'anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic>')

    def rect(self, p, x, y, w, h, color, dpi, name="Shape"):
        e = lambda v: max(int(round(v * EMU_PER_INCH / dpi)), 1)
        p._p.append(self._anchor(name, x, y, w, h, self._shape(e(w), e(h), color), dpi))

    def textbox(self, p, box, dpi):
        e = lambda v: max(int(round(v * EMU_PER_INCH / dpi)), 1)
        g = self._shape(e(box["w"]), e(box["h"]), None, box["xml"])
        p._p.append(self._anchor("Text", box["x"], box["y"], box["w"], box["h"], g, dpi))

    def picture(self, p, png_bytes, x, y, w, h, dpi):
        rid, _ = self.doc.part.get_or_add_image(io.BytesIO(png_bytes))
        e = lambda v: max(int(round(v * EMU_PER_INCH / dpi)), 1)
        g = ('<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
             '<pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="picture"/><pic:cNvPicPr/></pic:nvPicPr>'
             f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
             f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'
             '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
             '</a:graphicData></a:graphic>')
        p._p.append(self._anchor("Picture", x, y, w, h, g, dpi))


# ------------------------------------------------------------- pipeline
def process_page(img, dpi, a, tess):
    arr = np.array(img.convert("RGB"))
    H, W = arr.shape[:2]
    words = ocr_words(tess, img, a)
    bg = page_background(arr)
    check_ita = not getattr(a, "no_italic", False)
    for wd in words:
        wd["color"], wd["stroke"], wd["bg"], wd["italic"] = word_style(arr, wd, check_italic=check_ita)
    fills, images, lines, left = detect_graphics(arr, words, bg, a)
    if not a.no_sparse_pass:
        extra = orphan_words(tess, img, arr, left, images, a)
        for wd in extra:
            wd["color"], wd["stroke"], wd["bg"], wd["italic"] = word_style(arr, wd, check_italic=check_ita)
            wd["is_orphan"] = True
        words = words + extra
    if not getattr(a, "no_underline", False):
        lines = match_underlines(lines, words, a)
    if images and not a.keep_image_text:
        words = [d for d in words if not any(
            r["x"] <= d["x"] + d["w"] / 2 <= r["x"] + r["w"] and
            r["y"] <= d["y"] + d["h"] / 2 <= r["y"] + r["h"] for r in images)]
    boxes = make_text_boxes(words, dpi, a, page_w=W) if words else []
    crops = []
    for r in images:
        buf = io.BytesIO()
        Image.fromarray(arr[r["y"]:r["y"] + r["h"], r["x"]:r["x"] + r["w"]]).save(buf, "PNG")
        crops.append(buf.getvalue())
    return dict(w=W, h=H, bg=bg, fills=fills, images=images, crops=crops,
                lines=lines, boxes=boxes, arr=arr)


def add_page(doc, em, first, page, dpi):
    sec = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width = Emu(int(page["w"] / dpi * EMU_PER_INCH))
    sec.page_height = Emu(int(page["h"] / dpi * EMU_PER_INCH))
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin",
                 "header_distance", "footer_distance"):
        setattr(sec, attr, Emu(0))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)

    if sum(abs(c - 255) for c in page["bg"]) > 30:  # coloured paper
        em.rect(p, 0, 0, page["w"], page["h"], snap_color(page["bg"]), dpi, "Background")
    for r in page["fills"]:
        em.rect(p, r["x"], r["y"], r["w"], r["h"], r["color"], dpi, "Box")
    for r, png in zip(page["images"], page["crops"]):
        em.picture(p, png, r["x"], r["y"], r["w"], r["h"], dpi)
    for r in page["lines"]:
        em.rect(p, r["x"], r["y"], r["w"], r["h"], r["color"], dpi, "Line")
    for b in page["boxes"]:
        em.textbox(p, b, dpi)


def save_debug(page, path):
    img = cv2.cvtColor(page["arr"], cv2.COLOR_RGB2BGR)
    for r in page["fills"]:
        cv2.rectangle(img, (r["x"], r["y"]), (r["x"] + r["w"], r["y"] + r["h"]), (255, 0, 0), 4)
    for r in page["images"]:
        cv2.rectangle(img, (r["x"], r["y"]), (r["x"] + r["w"], r["y"] + r["h"]), (0, 0, 255), 6)
    for r in page["lines"]:
        cv2.rectangle(img, (r["x"], r["y"]), (r["x"] + r["w"], r["y"] + r["h"]), (0, 200, 0), 3)
    for b in page["boxes"]:
        cv2.rectangle(img, (int(b["x"]), int(b["y"])),
                      (int(b["x"] + b["w"]), int(b["y"] + b["h"])), (0, 140, 255), 2)
    cv2.imwrite(str(path), img)


def build_parser():
    p = argparse.ArgumentParser(
        description="Offline layout-preserving converter: image-based PDF/images -> editable DOCX.")
    p.add_argument("input", type=Path, help="PDF or image file")
    p.add_argument("-o", "--output", type=Path, help="Output .docx (default: <input>_layout.docx)")
    p.add_argument("--pages", help="Pages to convert, e.g. 1-3,5")
    p.add_argument("--dpi", type=int, default=300, help="Render DPI for PDFs (default 300)")
    p.add_argument("--lang", default="eng", help="Tesseract language(s), e.g. eng+hin")
    p.add_argument("--psm", type=int, help="Tesseract page segmentation mode")
    p.add_argument("--min-conf", type=float, default=40, help="Drop OCR words below this confidence")
    p.add_argument("--no-sparse-pass", action="store_true",
                   help="Skip the extra OCR pass that finds isolated numbers/labels (faster)")
    g = p.add_argument_group("Text look")
    g.add_argument("--font", default="Arial", help="Font for all text (default Arial)")
    g.add_argument("--font-scale", type=float, default=1.0, help="Multiply estimated font sizes")
    g.add_argument("--y-offset", type=float, default=0.0, help="Move all text down by N points")
    g.add_argument("--bold-factor", type=float, default=1.2,
                   help="Text this much thicker than the page's usual text is bold "
                        "(higher = less bold, default 1.2)")
    g.add_argument("--no-bold", action="store_true", help="Never mark text bold")
    g.add_argument("--no-italic", action="store_true", help="Never mark text italic")
    g.add_argument("--no-underline", action="store_true", help="Never convert lines to Word underlines")
    g.add_argument("--no-align", action="store_true", help="Never detect paragraph alignment (keep left)")
    g.add_argument("--box-per-line", action="store_true",
                   help="One text box per line instead of per paragraph (most exact positions)")
    g = p.add_argument_group("Pictures and shapes")
    g.add_argument("--min-image-pct", type=float, default=0.3,
                   help="Smallest picture, as %% of page area (default 0.3)")
    g.add_argument("--bg-tol", type=int, default=45, help="Colour tolerance vs paper colour")
    g.add_argument("--no-images", action="store_true", help="Do not extract pictures/filled boxes")
    g.add_argument("--no-lines", action="store_true", help="Do not draw box/table lines")
    g.add_argument("--keep-image-text", action="store_true",
                   help="Also create text for words found inside pictures")
    p.add_argument("--page-batch", type=int, default=5,
                   help="Number of PDF pages to render per Poppler batch (default 5)")
    p.add_argument("--image-page-width", type=float, default=8.27,
                   help="Page width in inches for image (non-PDF) inputs (default 8.27 = A4)")
    p.add_argument("--debug", action="store_true", help="Save *_debug.png showing what was detected")
    p.add_argument("--poppler-path")
    p.add_argument("--tesseract-cmd")
    return p


def main():
    a = build_parser().parse_args()
    if not a.input.exists():
        sys.exit(f"Not found: {a.input}")
    try:
        import pytesseract as tess
    except ImportError:
        sys.exit("pytesseract is not installed. Run: pip install -r requirements.txt")
    if a.tesseract_cmd:
        tess.pytesseract.tesseract_cmd = a.tesseract_cmd
    out = a.output or a.input.with_name(a.input.stem + "_layout.docx")
    try:
        doc = Document()
        em = Emitter(doc)
        for n, (img, dpi) in enumerate(load_pages(a.input, a), 1):
            print(f"Page {n}: reading ...", end="\r", flush=True)
            page = process_page(img, dpi, a, tess)
            add_page(doc, em, n == 1, page, dpi)
            if a.debug:
                save_debug(page, out.with_name(f"{out.stem}_p{n:03d}_debug.png"))
            print(f"Page {n}: {len(page['boxes'])} text boxes, {len(page['images'])} pictures, "
                  f"{len(page['fills'])} filled boxes, {len(page['lines'])} lines")
        doc.save(str(out))
        print(f"Saved: {out}")
    except SystemExit:
        raise
    except Exception as e:
        sys.exit(f"Error: {e}")


if __name__ == "__main__":
    main()
